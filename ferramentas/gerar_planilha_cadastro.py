# Planilha SIMPLES de cadastro de equipamentos (versão 2 — 23/10/2026).
# Quem preenche só informa Área, Setor e o nome do equipamento. O tipo (preventivas) e a
# categoria (chamado) são reconhecidos pelo nome, pela aba Reconhecimento.
# Uso: python3 gerar_planilha_cadastro.py [biblioteca.json] [saida.xlsx]
import json, sys, os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from reconhecimento import linhas as palavras

BIB = sys.argv[1] if len(sys.argv) > 1 else '/home/claude/biblioteca/biblioteca.json'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Cadastro_Equipamentos_RioMar_Recife.xlsx'
R = json.load(open(BIB))
NAVY = "13254A"; ORANGE = "E94D10"; GRAY = "5B6B85"; LIGHT = "EEF1F6"; INPUT = "FFF7E0"; GREEN = "1F7A45"
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="C9D3E3"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
FREQ = {"SEMANAL": "semanal", "MENSAL": "mensal", "TRIMESTRAL": "trimestral", "QUADRIMESTRAL": "quadrimestral", "SEMESTRAL": "semestral", "ANUAL": "anual", "QUINQUENAL": "a cada 5 anos"}

# Tipos → resumo das preventivas
tipos = collections.OrderedDict()
for r in R:
    tipos.setdefault(r["tipo_equipamento"], []).append(f'{FREQ[r["frequencia"]]}{" (terceirizada)" if r["executor"] == "TERCEIRIZADA" else ""}')
for t in ["Mobiliário", "TI e KDS", "Predial e estrutura", "Outro equipamento"]:
    tipos[t] = None

wb = Workbook()
# ------------------------------------------------------------------ Como preencher
ws = wb.active; ws.title = "Como preencher"; ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 6; ws.column_dimensions["C"].width = 100
ws["B2"] = "Cadastro de equipamentos — RioMar Recife"; ws["B2"].font = F(size=18, bold=True, color=NAVY)
ws["B3"] = "Plataforma de Manutenção Entre Amigos · planilha simplificada · 23/10/2026"; ws["B3"].font = F(size=10, color=GRAY)
passos = [
    ("1", "Vá na aba Equipamentos. Preencha uma linha por equipamento ou item, a partir da linha 6."),
    ("2", "Área e Setor: escolha na lista (ou digite um novo). Ex.: Cozinha › Sobremesa."),
    ("3", "Equipamento: escreva o nome como a equipe chama. Ex.: Fritadeira 01, Câmara de congelados, Cadeira, Tela KDS."),
    ("4", "Qtd: só se tiver vários iguais no mesmo setor (ex.: 12 cadeiras). O sistema cria Cadeira 01, 02… sozinho."),
    ("5", "Pronto. As colunas cinza mostram o que o sistema entendeu. Se aparecer \"Confira\", tudo bem: você escolhe o tipo na hora de importar."),
]
for i, (n, t) in enumerate(passos):
    a = ws.cell(row=5 + i * 2, column=2, value=n); a.font = F(size=16, bold=True, color="FFFFFF"); a.fill = fill(ORANGE)
    a.alignment = Alignment(horizontal="center", vertical="center")
    c = ws.cell(row=5 + i * 2, column=3, value=t); c.font = F(size=12, color=NAVY); c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5 + i * 2].height = 30
ws["B16"] = "Opcional"; ws["B16"].font = F(size=11, bold=True, color=NAVY)
ws["C16"] = ("Última preventiva grande: se souber, a data da última limpeza de duto, recarga de extintor, manutenção da câmara etc. "
             "O sistema calcula a próxima. Em branco, a preventiva entra já no primeiro mês.")
ws["C17"] = ("Temperaturas, tensão, fabricante e nº de série NÃO precisam ser preenchidos aqui. As faixas de temperatura "
             "têm um padrão na importação (congelados e resfriados) e podem ser ajustadas depois na ficha de cada equipamento.")
for c in ("C16", "C17"):
    ws[c].font = F(size=11, color=GRAY); ws[c].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[16].height = 46; ws.row_dimensions[17].height = 46

# ------------------------------------------------------------------ Áreas e setores
wa = wb.create_sheet("Áreas e setores"); wa.sheet_view.showGridLines = False
wa["A1"] = "Áreas e setores da loja"; wa["A1"].font = F(size=14, bold=True, color=NAVY)
wa["A2"] = "Sugestão inicial. Inclua ou renomeie à vontade: as listas da aba Equipamentos vêm daqui."; wa["A2"].font = F(size=10, color=GRAY)
for col, h in zip("AB", ["Área", "Setor"]):
    c = wa[f"{col}4"]; c.value = h; c.font = F(bold=True, color="FFFFFF"); c.fill = fill(NAVY); c.border = border
areas = [("Cozinha", "Cozinha quente"), ("Cozinha", "Cozinha fria"), ("Cozinha", "Sobremesa"), ("Cozinha", "Copa"), ("Cozinha", "Lavagem"),
         ("Estoque", "Câmaras frias"), ("Estoque", "Estoque seco"), ("Salão", "Salão principal"), ("Salão", "Salão frente do bar"), ("Salão", "Varanda"),
         ("Bar", "Balcão do bar"), ("Banheiros", "Banheiro clientes"), ("Banheiros", "Banheiro funcionários"), ("Área técnica", "Casa de máquinas"),
         ("Área técnica", "Quadros elétricos"), ("Administrativo", "Escritório"), ("Área externa", "Fachada")]
for i in range(80):
    for j, col in enumerate("AB"):
        c = wa[f"{col}{5 + i}"]; c.value = areas[i][j] if i < len(areas) else None; c.fill = fill(INPUT); c.border = border; c.font = F(size=11)
wa["D4"] = "Áreas (lista)"; wa["D4"].font = F(bold=True, color="FFFFFF"); wa["D4"].fill = fill(NAVY)
uni = list(dict.fromkeys(a for a, _ in areas))
for i in range(30):
    c = wa[f"D{5 + i}"]; c.value = uni[i] if i < len(uni) else None; c.fill = fill(INPUT); c.font = F(size=11)
for col, w in zip("ABD", [24, 30, 24]): wa.column_dimensions[col].width = w

# ------------------------------------------------------------------ Reconhecimento
wr = wb.create_sheet("Reconhecimento"); wr.sheet_view.showGridLines = False
wr["A1"] = "Como o sistema reconhece o equipamento pelo nome"; wr["A1"].font = F(size=14, bold=True, color=NAVY)
wr["A2"] = ("Se o nome tiver a palavra da coluna A, o equipamento recebe o Tipo e a Categoria da linha. Vale a última palavra da lista "
            "encontrada no nome. Pode incluir palavras novas no fim (ex.: um apelido que a equipe usa).")
wr["A2"].font = F(size=10, color=GRAY); wr["A2"].alignment = Alignment(wrap_text=True); wr.merge_cells("A2:C2"); wr.row_dimensions[2].height = 30
for j, h in enumerate(["Palavra no nome", "Tipo (define as preventivas)", "Categoria (chamado)"], 1):
    c = wr.cell(row=4, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = fill(NAVY); c.border = border
PAL = palavras(); RLAST = 4 + len(PAL) + 60
for i in range(len(PAL) + 60):
    vals = PAL[i] if i < len(PAL) else (None, None, None)
    for j, v in enumerate(vals, 1):
        c = wr.cell(row=5 + i, column=j, value=v); c.font = F(size=10); c.border = border
        if i >= len(PAL): c.fill = fill(INPUT)
for col, w in zip("ABC", [26, 38, 26]): wr.column_dimensions[col].width = w

# ------------------------------------------------------------------ Tipos (referência)
wt = wb.create_sheet("Tipos"); wt.sheet_view.showGridLines = False
wt["A1"] = "Tipos e as preventivas que cada um gera"; wt["A1"].font = F(size=14, bold=True, color=NAVY)
for j, h in enumerate(["Tipo", "Preventivas"], 1):
    c = wt.cell(row=4, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = fill(NAVY); c.border = border
row = 5
for t, v in tipos.items():
    wt.cell(row=row, column=1, value=t).font = F(size=10)
    wt.cell(row=row, column=2, value="; ".join(dict.fromkeys(v)).capitalize() if v else "Sem preventiva (só chamados)").font = F(size=10, italic=not v)
    row += 1
TLAST = row - 1
wt.column_dimensions["A"].width = 38; wt.column_dimensions["B"].width = 60
dv_tipo_lista = f"Tipos!$A$5:$A${TLAST}"

# ------------------------------------------------------------------ Equipamentos
we = wb.create_sheet("Equipamentos", 1); we.sheet_view.showGridLines = False; we.freeze_panes = "D4"
we["A1"] = "Equipamentos — RioMar Recife"; we["A1"].font = F(size=16, bold=True, color=NAVY)
we["A2"] = "Preencha a partir da linha 6 (linhas 4 e 5 são exemplos). Laranja = preencher · Azul = opcional · Cinza = o sistema preenche."
we["A2"].font = F(size=10, color=GRAY)
FIRST, LAST, EX = 4, 403, 6
cols = [("Área", ORANGE, 18), ("Setor", ORANGE, 24), ("Equipamento / item", ORANGE, 34), ("Qtd", NAVY, 7),
        ("Última preventiva grande", NAVY, 14), ("Observação", NAVY, 28),
        ("O sistema vai entender como", GRAY, 30), ("Categoria no chamado", GRAY, 22), ("Preventivas", GRAY, 30), ("Conferência", GRAY, 24)]
for j, (h, cor, w) in enumerate(cols, 1):
    c = we.cell(row=3, column=j, value=h); c.font = F(bold=True, color="FFFFFF", size=10); c.fill = fill(cor)
    c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    we.column_dimensions[c.column_letter].width = w
we.row_dimensions[3].height = 34
RK = f"Reconhecimento!$A$5:$A${RLAST}"
for r in range(FIRST, LAST + 1):
    for j in range(1, 7):
        c = we.cell(row=r, column=j); c.border = border; c.fill = fill(INPUT); c.font = F(size=10, italic=r < EX, color=GRAY if r < EX else "0F1B33")
    we.cell(row=r, column=5).number_format = "DD/MM/YYYY"
    achou = f'1/(({RK}<>"")*ISNUMBER(SEARCH(" "&{RK}," "&$C{r})))'
    we.cell(row=r, column=7, value=f'=IF($C{r}="","",IFERROR(LOOKUP(2,{achou},Reconhecimento!$B$5:$B${RLAST}),"Confira: escolha na importação"))')
    we.cell(row=r, column=8, value=f'=IF($C{r}="","",IFERROR(LOOKUP(2,{achou},Reconhecimento!$C$5:$C${RLAST}),"—"))')
    we.cell(row=r, column=9, value=f'=IF($C{r}="","",IFERROR(VLOOKUP($G{r},Tipos!$A$5:$B${TLAST},2,FALSE),"—"))')
    we.cell(row=r, column=10, value=(f'=IF(COUNTA($A{r}:$C{r})=0,"",IF($A{r}="","Falta a Área",IF($B{r}="","Falta o Setor",IF($C{r}="","Falta o nome",'
                                      f'IF(COUNTIF($C${EX}:$C${LAST},$C{r})>1,"Nome repetido: use 01, 02… ou a Qtd",'
                                      f'IF(LEFT($G{r},7)="Confira","Confira o tipo na importação","OK"))))))'))
    for j in range(7, 11):
        c = we.cell(row=r, column=j); c.fill = fill(LIGHT); c.border = border; c.font = F(size=10, color=GRAY); c.alignment = Alignment(wrap_text=True, vertical="top")
ex = [("Cozinha", "Sobremesa", "Micro-ondas", None, None, "Exemplo: item só para chamados"),
      ("Estoque", "Câmaras frias", "Câmara de congelados", None, None, "Exemplo: com preventiva")]
for i, vals in enumerate(ex):
    for j, v in enumerate(vals, 1):
        we.cell(row=FIRST + i, column=j, value=v)
we.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(formula=[f'$J{FIRST}="OK"'], font=Font(name="Arial", color=GREEN, bold=True), fill=fill("DDF3E4")))
we.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(formula=[f'LEFT($J{FIRST},7)="Confira"'], font=Font(name="Arial", color="8A5A00", bold=True), fill=fill("FFF1C9")))
we.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(formula=[f'AND($J{FIRST}<>"",$J{FIRST}<>"OK",LEFT($J{FIRST},7)<>"Confira")'], font=Font(name="Arial", color="9F1D12", bold=True), fill=fill("FDE3DF")))
we.conditional_formatting.add(f"G{FIRST}:G{LAST}", FormulaRule(formula=[f'LEFT($G{FIRST},7)="Confira"'], font=Font(name="Arial", color="8A5A00", bold=True)))
# Listas (aceitam valor novo: aviso em vez de bloqueio)
dvs = [
    (DataValidation(type="list", formula1="='Áreas e setores'!$D$5:$D$34", allow_blank=True, showErrorMessage=False), f"A{FIRST}:A{LAST}"),
    (DataValidation(type="list", formula1="='Áreas e setores'!$B$5:$B$84", allow_blank=True, showErrorMessage=False), f"B{FIRST}:B{LAST}"),
    (DataValidation(type="whole", operator="between", formula1="1", formula2="100", allow_blank=True, showErrorMessage=True, error="Quantidade de 1 a 100."), f"D{FIRST}:D{LAST}"),
    (DataValidation(type="date", operator="between", formula1="DATE(2015,1,1)", formula2="TODAY()", allow_blank=True, showErrorMessage=True, error="Digite a data (dd/mm/aaaa), até hoje."), f"E{FIRST}:E{LAST}"),
]
for dv, rng in dvs:
    we.add_data_validation(dv); dv.add(rng)
we.auto_filter.ref = f"A3:J{LAST}"
# Contador no topo
we["G1"] = "Preenchidas"; we["H1"] = "OK"; we["I1"] = "Conferir tipo"
we["G2"] = f'=COUNTA(C{EX}:C{LAST})'; we["H2"] = f'=COUNTIF(J{EX}:J{LAST},"OK")'; we["I2"] = f'=COUNTIF(J{EX}:J{LAST},"Confira*")'
for c in ("G1", "H1", "I1"): we[c].font = F(size=9, color=GRAY)
for c in ("G2", "H2", "I2"): we[c].font = F(size=14, bold=True, color=NAVY)

wb.save(OUT)
print("ok", OUT, len(PAL), "palavras", TLAST - 4, "tipos")
